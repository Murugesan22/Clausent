import os
from pathlib import Path
from typing import Dict, Tuple
from app.core.exceptions import ExtractionError
from app.core.logging_config import logger

class ExtractionResult:
    def __init__(self, full_text: str, page_texts: Dict[int, str], page_count: int, is_scanned: bool = False, message: str = ""):
        self.full_text = full_text
        self.page_texts = page_texts
        self.page_count = page_count
        self.is_scanned = is_scanned
        self.message = message

class ExtractionService:
    @staticmethod
    def extract_from_file(file_path: Path, file_type: str) -> ExtractionResult:
        """
        Extracts raw text and page references from PDF, DOCX, or TXT file.
        """
        if not file_path.exists():
            raise ExtractionError(f"Target file does not exist at {file_path}")

        file_type = file_type.lower().lstrip(".")
        if file_type == "pdf":
            return ExtractionService._extract_pdf(file_path)
        elif file_type == "docx":
            return ExtractionService._extract_docx(file_path)
        elif file_type == "txt":
            return ExtractionService._extract_txt(file_path)
        else:
            raise ExtractionError(f"Unsupported file format '.{file_type}'. Supported formats: PDF, DOCX, TXT.")

    @staticmethod
    def _extract_pdf(file_path: Path) -> ExtractionResult:
        page_texts: Dict[int, str] = {}
        full_text_parts = []
        is_scanned = False

        # 1. Try PyMuPDF (fitz)
        try:
            import fitz
            doc = fitz.open(str(file_path))
            if len(doc) == 0:
                raise ExtractionError("PDF document contains 0 pages.")

            for page_num in range(len(doc)):
                page = doc[page_num]
                text = page.get_text("text") or ""
                page_texts[page_num + 1] = text
                if text.strip():
                    full_text_parts.append(text)
            doc.close()

        except ImportError:
            # 2. Portable fallback to pypdf
            try:
                import pypdf
                reader = pypdf.PdfReader(str(file_path))
                if len(reader.pages) == 0:
                    raise ExtractionError("PDF document contains 0 pages.")

                for idx, page in enumerate(reader.pages):
                    text = page.extract_text() or ""
                    page_texts[idx + 1] = text
                    if text.strip():
                        full_text_parts.append(text)
            except Exception as e:
                logger.error(f"Failed to extract PDF using pypdf: {e}")
                raise ExtractionError(f"Error parsing PDF file: {str(e)}")
        except Exception as e:
            logger.error(f"Failed to extract PDF using fitz: {e}")
            raise ExtractionError(f"Error parsing PDF file: {str(e)}")

        full_text = "\n\n".join(full_text_parts).strip()

        if not full_text:
            is_scanned = True
            msg = (
                "No extractable text was found in the PDF. The document appears to be a scanned image "
                "or empty. OCR is currently required for image-only scans."
            )
            logger.warning(f"PDF {file_path.name} is image-only/scanned.")
            return ExtractionResult(full_text="", page_texts=page_texts, page_count=len(page_texts), is_scanned=True, message=msg)

        return ExtractionResult(full_text=full_text, page_texts=page_texts, page_count=len(page_texts), is_scanned=False)


    @staticmethod
    def _extract_docx(file_path: Path) -> ExtractionResult:
        try:
            import docx
            doc = docx.Document(str(file_path))
            paragraphs = []

            # Extract body paragraphs
            for p in doc.paragraphs:
                txt = p.text.strip()
                if txt:
                    paragraphs.append(txt)

            # Extract table cells if any
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        paragraphs.append(" | ".join(row_cells))

            full_text = "\n\n".join(paragraphs).strip()
            if not full_text:
                raise ExtractionError("The DOCX document contains no text content.")

            # DOCX does not have fixed physical page breaks until rendered
            page_texts = {1: full_text}
            return ExtractionResult(full_text=full_text, page_texts=page_texts, page_count=1)

        except Exception as e:
            logger.error(f"Failed to extract DOCX: {e}")
            raise ExtractionError(f"Error parsing DOCX file: {str(e)}")

    @staticmethod
    def _extract_txt(file_path: Path) -> ExtractionResult:
        encodings_to_try = ["utf-8", "utf-8-sig", "utf-16", "cp1252", "latin-1"]
        content = None

        for enc in encodings_to_try:
            try:
                with open(file_path, "r", encoding=enc) as f:
                    content = f.read()
                    break
            except (UnicodeDecodeError, UnicodeError):
                continue

        if content is None:
            raise ExtractionError("Unable to decode TXT file with standard encodings (UTF-8, UTF-16, Latin-1).")

        full_text = content.strip()
        if not full_text:
            raise ExtractionError("The text file is empty.")

        page_texts = {1: full_text}
        return ExtractionResult(full_text=full_text, page_texts=page_texts, page_count=1)
